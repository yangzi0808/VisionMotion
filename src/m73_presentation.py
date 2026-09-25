"""
VisionMotion —— M7.3A 最终项目展示材料生成脚本（PPTX）

本脚本把已经封板的项目结果整理为一份 10 页、16:9 的中文项目展示 PPT：

    docs/VisionMotion_Final_Presentation.pptx

设计原则（M7.3A）：
    - 所有科学数值直接取自封板结果（不重新计算、不重新读取视频、不引入新模型）；
    - 图表优先：图 A / 图 B / 图 C 直接嵌入已封板 PNG（不重新绘制、不拉伸）；
    - 明确外部公开数据性质、FFT 独立交叉验证语义、SHM 条件性参考语义；
    - 三层不确定度 A/B/C 分别列出，不合并；
    - 不使用“算法精度 / 真值 / 标准答案 / 准确率”等措辞；
    - 色盲友好配色（蓝 / 橙 / 绿 / 灰），与三张成果图语义一致。

依赖：python-pptx（仅本脚本使用的构建期依赖，不加入项目运行期 requirements.txt）。
PDF 导出：本脚本只生成 PPTX；如需 PDF，请使用本机 PowerPoint COM 导出（见项目报告说明）。

运行方式（在项目根目录下）：
    .venv\\Scripts\\python.exe src\\m73_presentation.py
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ============================================================
# 路径与全局常量
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]
RESULTS_DIR = PROJECT_ROOT / "results"
OUTPUT_PATH = PROJECT_ROOT / "docs" / "VisionMotion_Final_Presentation.pptx"

FIG_A = RESULTS_DIR / "EXP-EXT-LAB67-V1_fig_A_y0_timeseries.png"
FIG_B = RESULTS_DIR / "EXP-EXT-LAB67-V1_fig_B_frequency_comparison.png"
FIG_C = RESULTS_DIR / "EXP-EXT-LAB67-V1_fig_C_pipeline.png"

# 配色（Okabe–Ito 风格，色盲友好；与三张成果图语义一致）
BLUE = RGBColor(0x00, 0x72, 0xB2)
ORANGE = RGBColor(0xD5, 0x5E, 0x00)
GREEN = RGBColor(0x00, 0x9E, 0x73)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)
LIGHT_BLUE = RGBColor(0xE6, 0xF0, 0xF8)
LIGHT_ORANGE = RGBColor(0xFB, 0xE8, 0xD8)
LIGHT_GREEN = RGBColor(0xE2, 0xF4, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x1A)

FONT_NAME = "Microsoft YaHei"

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

TITLE_LEFT = Inches(0.6)
TITLE_TOP = Inches(0.32)
TITLE_WIDTH = Inches(12.1)
TITLE_HEIGHT = Inches(0.8)

BODY_TOP = Inches(1.42)
BODY_HEIGHT = Inches(5.25)

FOOTER_TOP = Inches(6.82)
FOOTER_HEIGHT = Inches(0.5)

EXTERNAL_FOOTER = "External public educational dataset — ComPADRE / OSP Lab67 Video 1（51.5 g）"

# 封板数值（只读引用，不得修改）
FROZEN = {
    "T_exp": "0.542593 s",
    "f_exp": "1.843003 Hz",
    "T_peak": "0.542593 s",
    "T_trough": "0.542593 s",
    "T_half": "0.542105 s",
    "f_fft": "1.875 Hz",
    "df_fft": "0.1875 Hz",
    "m": "0.0515 kg",
    "k": "7.05 N/m",
    "T_theory": "0.537018101 s",
    "f_theory": "1.862134625 Hz",
    "dT": "+0.005574899 s（+1.0381%）",
    "df": "−0.019131625 Hz（−1.0274%）",
    "rect_lo": "1.855281725 Hz",
    "rect_hi": "1.868975645 Hz",
    "unc_a": "A：m、k 不确定度传播尺度 ±0.3228%（频率 ±0.006011 Hz）",
    "unc_b": "B：实验统计标准误 SE_T ≈ 0.0038 s、SE_f ≈ 0.0129 Hz",
    "unc_c": "C：保守采样定位上界 ±0.0333 s、±0.1132 Hz（方法学上界，非统计误差）",
}

FORBIDDEN_WORDS = ("accuracy", "精度 1%", "算法精度", "真值", "标准答案", "算法准确率", "y0_px",
                   "工业级", "突破", "高精度测量系统")


# ============================================================
# 基础工具
# ============================================================


def set_run_font(run, size, bold=False, color=DARK, font=FONT_NAME):
    """
    设置字体（含中文字体 a:ea），保证中文不出现方框。
    """
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        element = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1])
        if element is None:
            from pptx.oxml.ns import qn

            element = rpr.makeelement(qn(tag), {})
            rpr.append(element)
        element.set("typeface", font)


def add_title(slide, text, subtitle=None):
    """
    统一页标题（+ 可选副标题）与标题下方强调线。
    """
    box = slide.shapes.add_textbox(TITLE_LEFT, TITLE_TOP, TITLE_WIDTH, TITLE_HEIGHT)
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    set_run_font(paragraph.add_run(), 30, bold=True, color=DARK)
    paragraph.runs[0].text = text
    if subtitle:
        sub = slide.shapes.add_textbox(TITLE_LEFT, TITLE_TOP + Inches(0.72), TITLE_WIDTH, Inches(0.34))
        sub_frame = sub.text_frame
        sub_frame.word_wrap = True
        sub_par = sub_frame.paragraphs[0]
        set_run_font(sub_par.add_run(), 14, color=GRAY)
        sub_par.runs[0].text = subtitle
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, TITLE_LEFT, TITLE_TOP + Inches(0.78),
                                  Inches(1.6), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()
    line.shadow.inherit = False


def add_footer(slide, page_number, note=EXTERNAL_FOOTER):
    """
    统一页脚：左侧页码，右侧外部数据声明 + 数据源说明。
    """
    left = slide.shapes.add_textbox(TITLE_LEFT, FOOTER_TOP, Inches(1.2), FOOTER_HEIGHT)
    left_frame = left.text_frame
    left_frame.word_wrap = False
    set_run_font(left_frame.paragraphs[0].add_run(), 11, color=GRAY)
    left_frame.paragraphs[0].runs[0].text = "%d / 10" % page_number

    right = slide.shapes.add_textbox(Inches(2.4), FOOTER_TOP, Inches(10.3), FOOTER_HEIGHT)
    right_frame = right.text_frame
    right_frame.word_wrap = True
    first = right_frame.paragraphs[0]
    first.alignment = PP_ALIGN.RIGHT
    set_run_font(first.add_run(), 10.5, color=GRAY)
    first.runs[0].text = note
    second = right_frame.add_paragraph()
    second.alignment = PP_ALIGN.RIGHT
    set_run_font(second.add_run(), 10.5, color=GRAY)
    second.runs[0].text = "数据源：results/EXP-EXT-LAB67-V1_trajectory.csv（frame 76–235，160 帧，30 fps）"


def add_textbox(slide, left, top, width, height, lines, align=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.TOP, space_after=6):
    """
    lines = [(text, size, bold, color), ...]
    """
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    for index, (text, size, bold, color) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = text
        set_run_font(run, size, bold=bold, color=color)
    return box


def add_card(slide, left, top, width, height, fill=LIGHT_GRAY, line_color=None, radius=0.04):
    """
    圆角卡片（用于分区、关键数字框）。
    """
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    if line_color is None:
        card.line.fill.background()
    else:
        card.line.color.rgb = line_color
        card.line.width = Pt(1.0)
    card.shadow.inherit = False
    try:
        card.adjustments[0] = radius
    except Exception:
        pass
    return card


def add_flow_chips(slide, labels, left, top, width, chip_height, fill, text_color, size=13):
    """
    横向流程条：若干圆角小卡片 + "→" 箭头。
    """
    count = len(labels)
    arrow_width = Inches(0.30)
    chip_width = Emu(int((width - arrow_width * (count - 1)) / count))
    x = left
    for index, label in enumerate(labels):
        chip = add_card(slide, x, top, chip_width, chip_height, fill=fill)
        frame = chip.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = label
        set_run_font(run, size, bold=True, color=text_color)
        x = Emu(int(x + chip_width))
        if index < count - 1:
            arrow = slide.shapes.add_textbox(x, top, arrow_width, chip_height)
            arrow_frame = arrow.text_frame
            arrow_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            arrow_par = arrow_frame.paragraphs[0]
            arrow_par.alignment = PP_ALIGN.CENTER
            arrow_run = arrow_par.add_run()
            arrow_run.text = "→"
            set_run_font(arrow_run, size + 3, bold=True, color=GRAY)
            x = Emu(int(x + arrow_width))
    return chip_width


def add_picture_fit(slide, image_path, left, top, max_width, max_height, center=True):
    """
    按原始宽高比嵌入图片（不拉伸），必要时在给定框内居中。
    """
    with Image.open(image_path) as image:
        pixel_width, pixel_height = image.size
    aspect = pixel_width / pixel_height
    width = max_width
    height = Emu(int(width / aspect))
    if height > max_height:
        height = max_height
        width = Emu(int(height * aspect))
    x = left
    y = top
    if center:
        x = Emu(int(left + (max_width - width) / 2))
        y = Emu(int(top + (max_height - height) / 2))
    slide.shapes.add_picture(str(image_path), x, y, width=width, height=height)
    return x, y, width, height, aspect, pixel_width, pixel_height


# ============================================================
# 十页内容
# ============================================================


def slide_01_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.28))
    band.fill.solid(); band.fill.fore_color.rgb = BLUE; band.line.fill.background(); band.shadow.inherit = False

    add_textbox(slide, Inches(1.0), Inches(1.55), Inches(11.3), Inches(1.2),
                [("VisionMotion", 54, True, BLUE)])
    add_textbox(slide, Inches(1.0), Inches(2.75), Inches(11.3), Inches(0.6),
                [("基于计算机视觉的低成本非接触式位移与振动监测", 22, True, DARK)])
    add_textbox(slide, Inches(1.0), Inches(3.42), Inches(11.3), Inches(0.5),
                [("从视频到周期、频率与物理参考的可复现测量流程", 15, False, GRAY)])

    add_flow_chips(slide, ["视 频", "y₀(t)", "T", "f", "FFT", "SHM 参考"],
                   Inches(1.0), Inches(4.35), Inches(11.3), Inches(0.62),
                   LIGHT_BLUE, BLUE, size=13)

    add_textbox(slide, Inches(1.0), Inches(5.35), Inches(11.3), Inches(0.9),
                [("外部公开教育数据 · 可解释的视觉测量流程", 14, False, GRAY),
                 ("图 / 数据：results/EXP-EXT-LAB67-V1_*、docs/M6.3_FINAL_REPORT.md", 12, False, GRAY)])
    add_footer(slide, 1)


def slide_02_why(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "为什么要从“视频”测量振动？")

    questions = [
        "普通视频能否提取稳定、可重复的位置特征？",
        "像素位置能否直接形成可分析的时间序列？",
        "如何从时间序列得到周期与频率？",
        "如何用独立证据检查测量结果？",
    ]
    lines = [("%d. %s" % (index + 1, text), 18, False, DARK) for index, text in enumerate(questions)]
    add_textbox(slide, TITLE_LEFT, BODY_TOP + Inches(0.35), Inches(7.1), Inches(3.6), lines, space_after=18)
    add_textbox(slide, TITLE_LEFT, Inches(5.35), Inches(7.1), Inches(0.7),
                [("重点不是深度学习，而是可解释、可复现的视觉测量流程。", 17, True, ORANGE)])

    add_card(slide, Inches(8.05), BODY_TOP + Inches(0.1), Inches(4.6), Inches(4.6), LIGHT_BLUE)
    add_textbox(slide, Inches(8.35), BODY_TOP + Inches(0.35), Inches(4.0), Inches(0.5),
                [("核心链路", 16, True, BLUE)])
    chain = ["视频", "→  y₀(t) 位置时间序列", "→  T 周期", "→  f 频率", "→  FFT 交叉验证", "→  SHM 条件性参考"]
    add_textbox(slide, Inches(8.35), BODY_TOP + Inches(0.95), Inches(4.0), Inches(3.4),
                [(item, 15, False, DARK) for item in chain], space_after=14)
    add_footer(slide, 2)


def slide_03_data(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "公开教学实验数据")

    badge = add_card(slide, TITLE_LEFT, BODY_TOP - Inches(0.05), Inches(3.9), Inches(0.5), ORANGE)
    badge_frame = badge.text_frame
    badge_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    badge_par = badge_frame.paragraphs[0]
    badge_par.alignment = PP_ALIGN.CENTER
    badge_run = badge_par.add_run()
    badge_run.text = "EXTERNAL / PUBLIC DATASET"
    set_run_font(badge_run, 14, bold=True, color=WHITE)

    add_card(slide, TITLE_LEFT, BODY_TOP + Inches(0.65), Inches(6.1), Inches(4.1), LIGHT_GRAY)
    add_textbox(slide, TITLE_LEFT + Inches(0.3), BODY_TOP + Inches(0.9), Inches(5.5), Inches(3.6), [
        ("来源：ComPADRE / Open Source Physics（OSP）", 16, False, DARK),
        ("条目 ID：16081", 16, True, BLUE),
        ("标题：Teaching Harmonic Motion Online Using Tracker", 14, False, DARK),
        ("使用资源：Lab67 Video 1", 16, True, BLUE),
        ("质量标签：51.5 g（= 0.0515 kg）", 16, False, DARK),
    ], space_after=14)

    add_card(slide, Inches(7.05), BODY_TOP + Inches(0.65), Inches(5.7), Inches(4.1), LIGHT_BLUE)
    add_textbox(slide, Inches(7.35), BODY_TOP + Inches(0.9), Inches(5.1), Inches(3.6), [
        ("视频参数：1920 × 1080，30 fps", 16, False, DARK),
        ("帧数：236 帧（总长约 7.87 s）", 16, False, DARK),
        ("正式分析区间：frame 76–235", 16, True, BLUE),
        ("160 帧，5.3333 s", 16, False, DARK),
        ("数据性质：External public educational dataset", 14, True, ORANGE),
        ("（外部公开教育数据，不是本项目自行采集）", 13, False, GRAY),
    ], space_after=16)
    add_footer(slide, 3)


def slide_04_method(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "从图像到位置时间序列")

    steps = ["ROI", "gray < 120", "暗连通域", "行宽带", "暗带上缘 y₀"]
    add_flow_chips(slide, steps, TITLE_LEFT, BODY_TOP + Inches(0.25), Inches(12.1), Inches(0.6),
                   LIGHT_BLUE, BLUE, size=13)

    add_card(slide, TITLE_LEFT, BODY_TOP + Inches(1.15), Inches(6.4), Inches(2.9), LIGHT_GRAY)
    add_textbox(slide, TITLE_LEFT + Inches(0.3), BODY_TOP + Inches(1.35), Inches(5.8), Inches(2.6), [
        ("冻结参数（M6.2-2B）", 15, True, GRAY),
        ("ROI：x ∈ [690, 780)，y ∈ [380, 760)", 14, False, DARK),
        ("灰度阈值：gray < 120；8 邻域，面积 ≥ 150 px", 14, False, DARK),
        ("行宽带：行跨度 ≥ 35 px，连续 ≥ 3 行", 14, False, DARK),
        ("几何条件：带宽 40–75 px，带高 ≤ 15 px", 14, False, DARK),
    ], space_after=10)

    add_card(slide, Inches(7.3), BODY_TOP + Inches(1.15), Inches(5.4), Inches(2.9), LIGHT_GREEN)
    add_textbox(slide, Inches(7.6), BODY_TOP + Inches(1.4), Inches(4.8), Inches(2.4), [
        ("追踪结果", 15, True, GRAY),
        ("160 / 160 frames detected", 26, True, GREEN),
        ("frame 76–235 连续无缺口", 14, False, DARK),
        ("frame 74–75 保持缺失（未放宽规则补值）", 13, False, GRAY),
    ], space_after=12)

    add_textbox(slide, TITLE_LEFT, Inches(6.0), Inches(12.1), Inches(0.6),
                [("正式位置特征选择为暗带上缘 y₀，不使用质心、面积或滤波预测器作为正式位置。", 15, True, ORANGE)])
    add_footer(slide, 4)


def slide_05_signal(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "原始视觉位置序列 y₀(t)")

    add_picture_fit(slide, FIG_A, TITLE_LEFT, BODY_TOP, Inches(8.55), Inches(5.0))
    add_textbox(slide, TITLE_LEFT, Inches(6.45), Inches(8.55), Inches(0.35),
                [("图 A：160 个原始采样点 + 10 个局部极大 + 10 个局部极小（未平滑、未滤波、未插值）", 11, False, GRAY)])

    add_card(slide, Inches(9.35), BODY_TOP, Inches(3.35), Inches(5.0), LIGHT_GRAY)
    add_textbox(slide, Inches(9.6), BODY_TOP + Inches(0.25), Inches(2.85), Inches(4.5), [
        ("实测范围", 15, True, GRAY),
        ("y₀ = 508 – 622 px", 19, True, BLUE),
        ("极值数量", 15, True, GRAY),
        ("10 峰 + 10 谷", 19, True, BLUE),
        ("数据处理", 15, True, GRAY),
        ("未平滑、未滤波、未插值", 14, False, DARK),
        ("时域结果", 15, True, GRAY),
        ("T_exp = %s" % FROZEN["T_exp"], 17, True, ORANGE),
        ("f_exp = %s" % FROZEN["f_exp"], 17, True, ORANGE),
    ], space_after=8)
    add_footer(slide, 5)


def slide_06_period_fft(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "时域测量与频域交叉验证")

    add_card(slide, TITLE_LEFT, BODY_TOP + Inches(0.1), Inches(6.0), Inches(4.3), LIGHT_ORANGE)
    add_textbox(slide, TITLE_LEFT + Inches(0.3), BODY_TOP + Inches(0.35), Inches(5.4), Inches(3.9), [
        ("时域（主结果）", 17, True, ORANGE),
        ("T_exp = %s" % FROZEN["T_exp"], 20, True, DARK),
        ("f_exp = %s" % FROZEN["f_exp"], 20, True, DARK),
        ("10 个峰 + 10 个谷", 15, False, DARK),
        ("峰—峰：%s" % FROZEN["T_peak"], 15, False, DARK),
        ("谷—谷：%s" % FROZEN["T_trough"], 15, False, DARK),
        ("半周期法：%s" % FROZEN["T_half"], 15, False, DARK),
    ], space_after=9)

    add_card(slide, Inches(6.9), BODY_TOP + Inches(0.1), Inches(5.8), Inches(4.3), LIGHT_GREEN)
    add_textbox(slide, Inches(7.2), BODY_TOP + Inches(0.35), Inches(5.2), Inches(3.9), [
        ("频域（独立交叉验证）", 17, True, GREEN),
        ("FFT 主峰：%s" % FROZEN["f_fft"], 20, True, DARK),
        ("频率分辨率 df：%s" % FROZEN["df_fft"], 16, False, DARK),
        ("仅对原始 y₀(t) 做 rFFT（仅减均值）", 14, False, GRAY),
        ("FFT 仅作独立交叉验证，不替代时域峰间周期测量。", 15, True, GREEN),
    ], space_after=12)
    add_footer(slide, 6)


def slide_07_shm(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "标准质量—弹簧 SHM 参考")

    add_card(slide, TITLE_LEFT, BODY_TOP + Inches(0.1), Inches(6.0), Inches(2.5), LIGHT_BLUE)
    add_textbox(slide, TITLE_LEFT + Inches(0.3), BODY_TOP + Inches(0.3), Inches(5.4), Inches(2.1), [
        ("T = 2π √(m / k)", 22, True, BLUE),
        ("f = 1 / (2π) √(k / m)", 22, True, BLUE),
        ("m = %s    k = %s" % (FROZEN["m"], FROZEN["k"]), 16, False, DARK),
    ], space_after=14)

    add_card(slide, Inches(6.9), BODY_TOP + Inches(0.1), Inches(5.8), Inches(2.5), LIGHT_GRAY)
    add_textbox(slide, Inches(7.2), BODY_TOP + Inches(0.35), Inches(5.2), Inches(2.1), [
        ("条件性参考值", 15, True, GRAY),
        ("T_theory = %s" % FROZEN["T_theory"], 19, True, ORANGE),
        ("f_theory = %s" % FROZEN["f_theory"], 19, True, ORANGE),
    ], space_after=12)

    add_textbox(slide, TITLE_LEFT, BODY_TOP + Inches(2.9), Inches(12.1), Inches(0.8),
                [("这是项目采用的标准模型参考计算，不是外部实验资料明文公式。", 17, True, ORANGE)])
    add_textbox(slide, TITLE_LEFT, BODY_TOP + Inches(3.65), Inches(12.1), Inches(1.1),
                [("外部资料未说明 0.0515 kg 是否包含全部挂具/夹持件质量，因此这里属于条件性的标准模型参考计算。", 13, False, GRAY),
                 ("本页未做有效质量修正，也未反演任何等效参数。", 13, False, GRAY)], space_after=6)
    add_footer(slide, 7)


def slide_08_difference(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "实验结果与模型参考的差异")

    add_picture_fit(slide, FIG_B, TITLE_LEFT, BODY_TOP, Inches(8.35), Inches(4.6))
    add_textbox(slide, TITLE_LEFT, Inches(6.12), Inches(8.35), Inches(0.35),
                [("图 B：实验测量（时域）/ 标准模型参考 / FFT 交叉验证的对照", 11, False, GRAY)])

    add_card(slide, Inches(9.15), BODY_TOP, Inches(3.55), Inches(4.6), LIGHT_GRAY)
    add_textbox(slide, Inches(9.4), BODY_TOP + Inches(0.2), Inches(3.05), Inches(4.2), [
        ("差异", 15, True, GRAY),
        ("ΔT = +0.005574899 s", 13, True, DARK),
        ("相对差异 +1.0381%", 12, False, GRAY),
        ("Δf = −0.019131625 Hz", 13, True, DARK),
        ("相对差异 −1.0274%", 12, False, GRAY),
        ("统一表述：约 1.0%", 17, True, ORANGE),
        ("参数矩形 f ∈ [%s, %s]" % (FROZEN["rect_lo"], FROZEN["rect_hi"]), 12, False, DARK),
        ("实验值 1.843003 Hz 位于参数矩形下界之外", 12, True, DARK),
        (FROZEN["unc_a"], 10.5, False, GRAY),
        (FROZEN["unc_b"], 10.5, False, GRAY),
        (FROZEN["unc_c"], 10.5, False, GRAY),
    ], space_after=7)

    add_textbox(slide, Inches(9.15), BODY_TOP + Inches(4.72), Inches(3.55), Inches(0.55),
                [("差异超过 m、k 不确定度传播尺度，但未超过当前设定的保守采样定位上界；具体物理归因未验证。", 11, False, GRAY)])
    add_footer(slide, 8)


def slide_09_pipeline(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "VisionMotion 完整技术链")

    add_picture_fit(slide, FIG_C, TITLE_LEFT, BODY_TOP, Inches(8.9), Inches(4.9))
    add_textbox(slide, TITLE_LEFT, Inches(6.35), Inches(8.9), Inches(0.35),
                [("图 C：计算机视觉 → 时域测量 → 频域交叉验证 → 物理模型", 11, False, GRAY)])

    cards = [
        ("视觉层", "可解释检测规则（ROI / 阈值 / 连通域 / 行宽带）", BLUE, LIGHT_BLUE),
        ("测量层", "160 帧正式轨迹；T、f", ORANGE, LIGHT_ORANGE),
        ("验证层", "FFT 交叉验证 + 标准 SHM 参考 + 不确定度边界", GREEN, LIGHT_GREEN),
    ]
    top = BODY_TOP + Inches(0.1)
    for title, detail, color, fill in cards:
        add_card(slide, Inches(9.7), top, Inches(3.0), Inches(1.35), fill)
        add_textbox(slide, Inches(9.9), top + Inches(0.15), Inches(2.6), Inches(1.05), [
            (title, 15, True, color),
            (detail, 11.5, False, DARK),
        ], space_after=4)
        top = Emu(int(top + Inches(1.55)))

    add_textbox(slide, Inches(9.7), BODY_TOP + Inches(4.75), Inches(3.0), Inches(0.8),
                [("当前阶段属于视觉测量与方法验证", 12, True, ORANGE)])
    add_footer(slide, 9)


def slide_10_summary(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "现阶段成果与下一步")

    add_card(slide, TITLE_LEFT, BODY_TOP + Inches(0.1), Inches(6.0), Inches(4.3), LIGHT_GREEN)
    add_textbox(slide, TITLE_LEFT + Inches(0.3), BODY_TOP + Inches(0.3), Inches(5.4), Inches(3.9), [
        ("已经完成", 17, True, GREEN),
        ("✓ 视频 → 位置时间序列", 15, False, DARK),
        ("✓ 160/160 全检出", 15, False, DARK),
        ("✓ 周期 / 频率（T、f）", 15, False, DARK),
        ("✓ FFT 交叉验证", 15, False, DARK),
        ("✓ 标准 SHM 条件性参考", 15, False, DARK),
        ("✓ 不确定度与证据边界", 15, False, DARK),
        ("✓ 最终可复现报告", 15, False, DARK),
    ], space_after=7)

    add_card(slide, Inches(6.9), BODY_TOP + Inches(0.1), Inches(5.8), Inches(4.3), LIGHT_GRAY)
    add_textbox(slide, Inches(7.2), BODY_TOP + Inches(0.3), Inches(5.2), Inches(3.9), [
        ("当前限制", 17, True, GRAY),
        ("· 外部公开数据，不是自主实验采集", 14, False, DARK),
        ("· 单一视频 / 单一质量（51.5 g）", 14, False, DARK),
        ("· 30 fps 限制极值时间定位", 14, False, DARK),
        ("· y₀ 仍为像素（px）", 14, False, DARK),
        ("· 本阶段没有有效动态 px/mm 标定", 14, False, DARK),
        ("· 无阻尼 / 有效质量独立测量", 14, False, DARK),
        ("· 具体物理差异原因未验证", 14, False, DARK),
    ], space_after=7)

    add_textbox(slide, TITLE_LEFT, Inches(6.05), Inches(12.1), Inches(0.75), [
        ("项目报告：docs/M6.3_FINAL_REPORT.md　　数据来源说明：docs/ATTRIBUTION.md", 13, False, DARK),
        ("GitHub repository: https://github.com/yangzi0808/VisionMotion", 13, True, GRAY),
    ], space_after=4)
    add_footer(slide, 10)


# ============================================================
# 自检
# ============================================================


def iterate_text(prs):
    """
    遍历所有幻灯片中的文字（含表格/卡片/页脚）。
    """
    for index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip():
                            yield index, run.text


def check_presentation(prs, output_path, image_checks):
    """
    结构自检：页数 / 16:9 / 图片宽高比 / 封板数值 / 关键声明 / 禁用词 / 文本是否越界。
    """
    checks = []

    def add(desc, ok, extra=""):
        checks.append((desc, bool(ok), extra))

    slides = list(prs.slides)
    add("页数 = %d（要求 10）" % len(slides), len(slides) == 10)
    width_emu, height_emu = prs.slide_width, prs.slide_height
    add("页面尺寸 %.3f × %.3f in（宽高比 %.4f，16:9 ≈ 1.7778）"
        % (width_emu / 914400, height_emu / 914400, width_emu / height_emu),
        abs(width_emu / height_emu - 16 / 9) < 0.01)

    for label, aspect_src, aspect_used, px, size in image_checks:
        add("%s 宽高比未拉伸（源 %.3f×%.3f px，插入比例差 %.4f%%）"
            % (label, px[0], px[1], abs(aspect_used - aspect_src) / aspect_src * 100),
            abs(aspect_used - aspect_src) / aspect_src < 0.005)

    all_text = " ".join(text for _, text in iterate_text(prs))
    for token, desc in [
        (FROZEN["T_exp"], "T_exp"), (FROZEN["f_exp"], "f_exp"), (FROZEN["f_fft"], "FFT 主峰"),
        (FROZEN["df_fft"], "FFT 分辨率"), (FROZEN["T_peak"], "峰—峰周期"), (FROZEN["T_trough"], "谷—谷周期"),
        (FROZEN["T_half"], "半周期法"), (FROZEN["m"], "m"), (FROZEN["k"], "k"),
        (FROZEN["T_theory"], "T_theory"), (FROZEN["f_theory"], "f_theory"),
        ("+0.005574899 s", "ΔT"), ("+1.0381%", "ΔT 相对"), ("−0.019131625 Hz", "Δf"), ("−1.0274%", "Δf 相对"),
        (FROZEN["rect_lo"], "参数矩形下界"), (FROZEN["rect_hi"], "参数矩形上界"),
    ]:
        add("封板数值存在：%s" % desc, token in all_text, "" if token in all_text else "缺失 %s" % token)

    for token, desc in [
        ("External public educational dataset", "外部数据声明"),
        ("不是本项目自行采集", "非自采声明"),
        ("FFT 仅作独立交叉验证，不替代时域峰间周期测量", "FFT 独立交叉验证语义"),
        ("这是项目采用的标准模型参考计算，不是外部实验资料明文公式", "SHM 条件性声明（模型来源）"),
        ("外部资料未说明 0.0515 kg 是否包含全部挂具/夹持件质量", "SHM 条件性声明（质量构成）"),
        ("具体物理归因未验证", "物理归因未验证"),
        ("本阶段没有有效动态 px/mm 标定", "px/mm 未完成表述"),
        ("GitHub repository: https://github.com/yangzi0808/VisionMotion", "GitHub 仓库链接"),
        ("A：", "不确定度 A 单独列出"), ("B：", "不确定度 B 单独列出"), ("C：", "不确定度 C 单独列出"),
    ]:
        add(desc, token in all_text, "" if token in all_text else "缺失 %s" % token)

    hits = [(w, text) for _, text in iterate_text(prs) for w in FORBIDDEN_WORDS if w in text]
    add("禁用措辞 0 命中（%d 项关键词）" % len(FORBIDDEN_WORDS), not hits, str(hits[:3]))

    overflow = []
    for index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.left is None or shape.top is None:
                continue
            left, top = shape.left, shape.top
            right = left + (shape.width or 0)
            bottom = top + (shape.height or 0)
            if left < -Emu(1) or top < -Emu(1) or right > prs.slide_width + Emu(1) or bottom > prs.slide_height + Emu(1):
                overflow.append((index, shape.shape_type, left / 914400, top / 914400,
                                 right / 914400, bottom / 914400))
    add("所有形状位于页面范围内（越界 %d 个）" % len(overflow), not overflow, str(overflow[:3]))
    return checks


# ============================================================
# 入口
# ============================================================


def build_presentation(output_path=OUTPUT_PATH):
    """
    生成 10 页 PPTX 并返回 (prs, checks)。
    """
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    prs.core_properties.title = "VisionMotion — 基于计算机视觉的低成本非接触式位移与振动监测"
    prs.core_properties.author = "VisionMotion 项目（作者简介见 docs/AUTHOR_PROFILE.md）"
    prs.core_properties.subject = "外部公开教育数据上的视觉测量、周期/频率与条件性物理参考"

    slide_01_cover(prs)
    slide_02_why(prs)
    slide_03_data(prs)
    slide_04_method(prs)
    slide_05_signal(prs)
    slide_06_period_fft(prs)
    slide_07_shm(prs)
    slide_08_difference(prs)
    slide_09_pipeline(prs)
    slide_10_summary(prs)

    image_checks = []
    for label, path in (("图 A", FIG_A), ("图 B", FIG_B), ("图 C", FIG_C)):
        with Image.open(path) as image:
            pixel_width, pixel_height = image.size
        aspect_src = pixel_width / pixel_height
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.shape_type == 13 and shape.width and shape.height:  # PICTURE
                    if abs((shape.width / shape.height) - aspect_src) < 0.01:
                        image_checks.append((label, aspect_src, shape.width / shape.height,
                                             (pixel_width, pixel_height),
                                             (shape.width, shape.height)))
                        break

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    checks = check_presentation(prs, output_path, image_checks)
    return prs, checks


def main():
    """
    生成 PPTX，打印自检结果，返回退出码。
    """
    print("")
    print("========== M7.3A 最终项目展示材料（PPTX） ==========")
    print("输出：%s" % OUTPUT_PATH)
    for label, path in (("图 A", FIG_A), ("图 B", FIG_B), ("图 C", FIG_C)):
        print("嵌入 %s：%s（存在=%s）" % (label, path.name, path.exists()))

    prs, checks = build_presentation()
    print("")
    print("========== 自检明细 ==========")
    for index, (desc, ok, extra) in enumerate(checks, start=1):
        print("%2d. [%s] %s%s" % (index, "通过" if ok else "未通过", desc, ("  %s" % extra) if extra else ""))
    passed = all(ok for _, ok, _ in checks)
    print("自检结论：%s" % ("全部通过" if passed else "存在未通过项"))
    print("PPTX 体积：%d 字节" % OUTPUT_PATH.stat().st_size)
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
